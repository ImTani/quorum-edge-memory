"""Generate the Quorum pitch deck (16:9, dark theme) with python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import os

BG = RGBColor(0x0A, 0x0E, 0x16)
PANEL = RGBColor(0x13, 0x19, 0x25)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x9A, 0xA4, 0xB2)
AMBER = RGBColor(0xF5, 0xA5, 0x24)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
RED = RGBColor(0xEF, 0x44, 0x44)
INDIGO = RGBColor(0x7C, 0x8C, 0xFF)
FONT = "Segoe UI"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(num=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if num:
        text(s, 12.2, 7.0, 0.9, 0.35, f"{num} / 10", 11, MUTED, align=PP_ALIGN.RIGHT)
        text(s, 0.6, 7.0, 4, 0.35, "QUORUM", 11, MUTED, bold=True)
    return s


def text(s, x, y, w, h, content, size=18, color=WHITE, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """content: str or list of str / (str, dict) runs-per-paragraph."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        opts = {}
        if isinstance(line, tuple):
            line, opts = line
        r = p.add_run()
        r.text = line
        f = r.font
        f.name = opts.get("font", font)
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", italic)
        f.color.rgb = opts.get("color", color)
    return tb


def box(s, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        content=None, size=14, color=WHITE, bold=False, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, font=FONT, lw=1.5):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.08
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    if content is not None:
        lines = content if isinstance(content, list) else [content]
        for i, line_ in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            opts = {}
            if isinstance(line_, tuple):
                line_, opts = line_
            r = p.add_run()
            r.text = line_
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = opts.get("color", color)
    return sh


def line(s, x1, y1, x2, y2, color=MUTED, width=2, dash=False, arrow_end=True, arrow_start=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if arrow_start:
        h = etree.SubElement(ln, qn("a:headEnd"))
        h.set("type", "triangle")
    if arrow_end:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
    return c


def dot(s, x, y, d, color):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    c.line.fill.background()
    return c


def header(s, kicker, title, kcolor=INDIGO):
    text(s, 0.6, 0.45, 12, 0.4, kicker.upper(), 13, kcolor, bold=True)
    text(s, 0.6, 0.8, 12.1, 0.8, title, 32, WHITE, bold=True)


# ---------------------------------------------------------------- 1. Title
s = new_slide()
for i, c in enumerate([AMBER, GREEN, RED, INDIGO]):
    dot(s, 0.8 + i * 0.35, 1.3, 0.2, c)
text(s, 0.7, 1.7, 12, 1.3, "QUORUM", 72, WHITE, bold=True)
text(s, 0.75, 2.95, 12, 0.7, "Offline-first team memory on Qdrant Edge", 28, MUTED)
box(s, 0.8, 3.85, 0.08, 1.1, fill=AMBER, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.05, 3.8, 11, 1.2, [
    ("Code Cubicle 6.0", {"size": 20, "bold": True}),
    ("PS3: AI-Powered Edge Memory & Intelligence", {"size": 18, "color": MUTED}),
], spacing=4)
text(s, 0.75, 5.6, 12, 0.5, "Tanishk  ·  Tushar  ·  Aayat  ·  Lakshya", 18, WHITE)
text(s, 0.75, 6.1, 12, 0.5, "github.com/ImTani/quorum-edge-memory", 16, INDIGO, font=MONO)

# ---------------------------------------------------------------- 2. Problem
s = new_slide(2)
header(s, "The problem", "One deadline, two versions, no signal", RED)
text(s, 0.6, 1.85, 5.6, 4.5, [
    "Small field teams: film crews, event teams, agencies, site crews.",
    "Many clients at once, across email, WhatsApp and conversations on set.",
    "Often somewhere with no signal at all.",
    "Nobody notices the mismatch until it's too late.",
    ("Cloud assistants need a connection, and would have to read everything. That data shouldn't leave the device.",
     {"color": MUTED, "size": 16}),
], 19, WHITE, spacing=12)
# Sharma example
box(s, 6.8, 1.9, 5.9, 1.35, fill=PANEL, line=INDIGO, align=PP_ALIGN.LEFT, content=[
    ("EMAIL  ·  from client:sharma", {"size": 12, "color": INDIGO, "bold": True}),
    ("\"Final delivery on the 16th.\"", {"size": 20}),
])
box(s, 6.8, 3.55, 5.9, 1.35, fill=PANEL, line=AMBER, align=PP_ALIGN.LEFT, content=[
    ("ON SET  ·  no signal  ·  Lakshya", {"size": 12, "color": AMBER, "bold": True}),
    ("\"Sharma delivery is the 18th.\"", {"size": 20}),
])
box(s, 6.8, 5.2, 5.9, 0.9, fill=RGBColor(0x2A, 0x12, 0x14), line=RED,
    content=[("Sharma delivery: 16th vs 18th. Which one is true?", {"size": 17, "bold": True, "color": RED})])

# ---------------------------------------------------------------- 3. Solution
s = new_slide(3)
header(s, "What Quorum is", "A shared, offline-first memory layer for small teams", GREEN)
cards = [
    (AMBER, "Every laptop is a full edge node", "Hybrid search, extraction and answers run on-device. Works with zero signal."),
    (GREEN, "Syncs when it reconnects", "A persistent outbox survives restarts and drains to the team hub."),
    (RED, "Never silently picks a winner", "When teammates' facts disagree, Qdrant surfaces the conflict and the task owner resolves it."),
]
for i, (c, t, d) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 1.95, 3.85, 1.9, fill=PANEL, line=c, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, content=[
        (t, {"size": 19, "bold": True, "color": c}),
        (d, {"size": 15, "color": MUTED}),
    ])
text(s, 0.6, 4.4, 12.1, 1.8, [
    ("Your team's memory, on every laptop, even with no signal.", {"bold": True}),
    "It reads everything, and nothing leaves unless it's work your team needs.",
    "When two people disagree about a deadline, it asks instead of guessing.",
], 20, WHITE, spacing=6)

# ---------------------------------------------------------------- 4. Architecture
s = new_slide(4)
header(s, "Architecture", "Two edge laptops, one home hub, sync when online")
stack = [
    "Qdrant Edge (vector memory)",
    "FastEmbed: dense + BM25 sparse",
    "SQLite: graph + outbox",
    "Local LLM (llama.cpp / Ollama)",
]


def laptop(x, y, name):
    box(s, x, y, 3.6, 3.3, fill=PANEL, line=AMBER, lw=2)
    text(s, x + 0.15, y + 0.1, 3.3, 0.45, name, 18, AMBER, bold=True)
    for i, t in enumerate(stack):
        box(s, x + 0.2, y + 0.65 + i * 0.63, 3.2, 0.5, fill=BG, line=None, content=t, size=13, color=WHITE)


laptop(0.6, 1.9, "Laptop A  ·  edge node")
laptop(9.1, 1.9, "Laptop B  ·  edge node")
box(s, 4.95, 2.65, 3.4, 1.8, fill=PANEL, line=GREEN, lw=2, content=[
    ("Home hub", {"size": 18, "bold": True, "color": GREEN}),
    ("Qdrant Server (Docker)", {"size": 14}),
    ("my-devices + team tiers only", {"size": 12, "color": MUTED}),
])
line(s, 4.2, 3.55, 4.95, 3.55, GREEN, 2.5, arrow_start=True)
line(s, 8.35, 3.55, 9.1, 3.55, GREEN, 2.5, arrow_start=True)
text(s, 4.95, 2.2, 3.4, 0.4, "sync when online", 13, GREEN, align=PP_ALIGN.CENTER)
# inputs
box(s, 0.6, 5.75, 1.6, 0.6, fill=PANEL, line=INDIGO, content="Gmail", size=14)
box(s, 2.35, 5.75, 1.85, 0.6, fill=PANEL, line=INDIGO, content="WhatsApp export", size=13)
line(s, 1.4, 5.75, 1.4, 5.2, INDIGO, 2)
line(s, 3.27, 5.75, 3.27, 5.2, INDIGO, 2)
text(s, 0.6, 6.4, 3.8, 0.4, "read-only, cached locally", 12, MUTED)
box(s, 5.4, 5.75, 2.5, 0.6, fill=None, line=MUTED, content="Phone (mocked)", size=14, color=MUTED)
line(s, 5.4, 6.05, 4.2, 4.9, MUTED, 1.5, dash=True)
text(s, 8.6, 5.75, 4.2, 1.0, [
    ("Edge app shell: Tauri (Rust + TypeScript)", {"size": 13, "color": MUTED}),
    ("Private path makes no cloud calls", {"size": 13, "color": AMBER}),
], spacing=4)

# ---------------------------------------------------------------- 5. Memory model
s = new_slide(5)
header(s, "Memory model", "Claims, not facts")
claim = """{
  "entity_id": "task_sharma_edit",
  "attribute": "due_date",
  "value": "2026-10-16",
  "source": {"kind": "email",
             "author": "client:sharma"},
  "stated_at": "2026-10-03T11:20+05:30",
  "captured_by": "lakshya",
  "tier": "team",
  "status": "active",
  "conflict_id": null,
  "version": 1
}"""
box(s, 0.6, 1.85, 6.3, 4.85, fill=PANEL, line=INDIGO, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT,
    content=[(l, {"font": MONO, "size": 16, "color": WHITE}) for l in claim.split("\n")])
text(s, 7.3, 1.85, 5.4, 1.5, [
    ("Every fact is stored with its source.", {"bold": True}),
    ("Who said it, where, when, on which device.", {"color": MUTED, "size": 16}),
], 19, spacing=4)
statuses = [
    (GREEN, "active", "current accepted value"),
    (RED, "disputed", "open conflict, both shown"),
    (MUTED, "superseded", "kept for history"),
    (INDIGO, "retracted", "withdrawn by its author"),
]
for i, (c, n, d) in enumerate(statuses):
    y = 3.3 + i * 0.62
    dot(s, 7.35, y + 0.13, 0.24, c)
    text(s, 7.75, y, 1.8, 0.5, n, 17, c, bold=True, font=MONO)
    text(s, 9.6, y + 0.03, 3.2, 0.5, d, 15, MUTED)
text(s, 7.3, 5.9, 5.4, 0.8, "Graph in SQLite: Person, Project, Task, Event, Source. Vectors in Qdrant Edge: dense + BM25 per claim.",
     13, MUTED)

# ---------------------------------------------------------------- 6. Qdrant's job
s = new_slide(6)
header(s, "Qdrant's job", "Qdrant does the conflict detection and enforces privacy")
# pipeline
steps = [
    (AMBER, "New claim", "Sharma delivery = 18th"),
    (INDIGO, "Filtered hybrid search", "dense + BM25, filter: attribute = due_date, status = active"),
    (RED, "Close match, different value", "Conflict opened"),
]
for i, (c, t, d) in enumerate(steps):
    x = 0.6 + i * 4.15
    box(s, x, 1.9, 3.7, 1.55, fill=PANEL, line=c, content=[
        (t, {"size": 17, "bold": True, "color": c}),
        (d, {"size": 13, "color": MUTED}),
    ])
    if i < 2:
        line(s, x + 3.7, 2.67, x + 4.15, 2.67, MUTED, 2)
box(s, 0.6, 3.75, 12.1, 0.95, fill=RGBColor(0x2A, 0x12, 0x14), line=RED, content=[
    ("\"Found 2 claims about the Sharma delivery, similarity 0.91. Dates disagree.\"",
     {"size": 20, "italic": True, "color": WHITE}),
])
box(s, 0.6, 5.0, 12.1, 1.6, fill=PANEL, line=GREEN, align=PP_ALIGN.LEFT, content=[
    ("Tier filter on every search", {"size": 18, "bold": True, "color": GREEN}),
    ("Every Qdrant query carries a payload filter on tier, so shared views can never surface device-only claims.",
     {"size": 15, "color": MUTED}),
    ("Detected at ingest, and at sync when claims from separate offline devices meet for the first time.",
     {"size": 15, "color": MUTED}),
])

# ---------------------------------------------------------------- 7. Conflict flow
s = new_slide(7)
header(s, "Conflict flow", "The owner decides. Nothing is deleted.", RED)
flow = [
    (RED, "Detected", "at ingest or at sync"),
    (AMBER, "Open", "both claims disputed; answers show both + sources"),
    (INDIGO, "Owner resolves", "picks or corrects"),
    (GREEN, "Resolved", "chosen claim active, other superseded (kept)"),
]
xs = [0.6, 3.75, 6.9, 10.05]
for (c, t, d), x in zip(flow, xs):
    box(s, x, 1.95, 2.65, 1.7, fill=PANEL, line=c, lw=2, content=[
        (t, {"size": 19, "bold": True, "color": c}),
        (d, {"size": 13, "color": MUTED}),
    ])
for x in xs[:3]:
    line(s, x + 2.65, 2.8, x + 3.15, 2.8, MUTED, 2)
# ask teammate branch
box(s, 5.3, 4.35, 3.4, 1.25, fill=PANEL, line=INDIGO, content=[
    ("Ask teammate", {"size": 17, "bold": True, "color": INDIGO}),
    ("\"Should I ask Lakshya?\" Drafts it, sends only on yes", {"size": 13, "color": MUTED}),
])
line(s, 5.05, 3.65, 5.9, 4.35, MUTED, 1.5, dash=True)
text(s, 2.6, 3.85, 2.3, 0.4, "owner hasn't answered", 12, MUTED, align=PP_ALIGN.RIGHT)
line(s, 8.7, 4.95, 11.35, 3.65, MUTED, 1.5, dash=True)
text(s, 9.9, 4.55, 2.9, 0.4, "reply confirms a value", 12, MUTED)
box(s, 0.6, 6.0, 12.1, 0.75, fill=None, line=AMBER, content=[
    ("\"Disputed. The client's email says the 16th; you said the 18th on set.\"",
     {"size": 17, "italic": True, "color": AMBER}),
])

# ---------------------------------------------------------------- 8. Sync tiers + privacy
s = new_slide(8)
header(s, "Sync tiers + privacy", "Everything starts device-only", AMBER)
tiers = [
    (AMBER, "Device only", "This device", "Nowhere", "Personal finances, window titles"),
    (INDIGO, "My devices", "My devices + hub", "My other devices", "Personal schedule, birthdays"),
    (GREEN, "Team", "All team devices + hub", "Whole team", "Client deadlines, shoot schedules"),
]
cols = [("Tier", 0.6, 2.3), ("Stored on", 2.9, 3.0), ("Syncs to", 5.9, 2.6), ("Examples", 8.5, 4.2)]
for name, x, w in cols:
    text(s, x, 1.8, w, 0.4, name.upper(), 12, MUTED, bold=True)
for i, (c, *vals) in enumerate(tiers):
    y = 2.25 + i * 0.72
    box(s, 0.6, y, 12.1, 0.62, fill=PANEL, line=None)
    box(s, 0.6, y, 0.08, 0.62, fill=c, shape=MSO_SHAPE.RECTANGLE)
    for (name, x, w), v in zip(cols, vals):
        text(s, x + 0.15, y + 0.1, w - 0.15, 0.45, v, 15, c if name == "Tier" else WHITE,
             bold=(name == "Tier"))
text(s, 0.6, 4.5, 12.1, 0.5,
     "A claim moves up a tier only when the local model classifies it that way. Unsure? It stays on the device.",
     15, MUTED)
priv = [
    (AMBER, "Persistent outbox", "Dual write + background worker with retries; survives restarts"),
    (GREEN, "Partial snapshots", "Down-sync sends only changed segments"),
    (INDIGO, "No cloud calls", "Extraction, classification and answers run on the local model"),
    (RED, "Encrypted at rest", "Memory, graph and outbox"),
]
for i, (c, t, d) in enumerate(priv):
    x = 0.6 + i * 3.05
    box(s, x, 5.15, 2.85, 1.25, fill=PANEL, line=c, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, content=[
        (t, {"size": 15, "bold": True, "color": c}),
        (d, {"size": 12, "color": MUTED}),
    ])

# ---------------------------------------------------------------- 9. Demo plan + real vs mocked
s = new_slide(9)
header(s, "Demo plan", "Two physical laptops: one offline, one online")
demo = [
    "1.  Tanishk's laptop goes offline (airplane mode, network monitor at 0 B/s). Search still answers on-device.",
    "2.  Offline, Tanishk types a quick note: \"Client says push the Sharma delivery to the 18th\". It shows amber, outbox 1.",
    "3.  Meanwhile Lakshya's laptop, still online, gets the client email: \"delivery confirmed for the 16th\".",
    "4.  Tanishk reconnects, the outbox drains, and Qdrant flags \"2 claims, similarity 0.91, dates disagree\".",
    "5.  The owner is asked, no guessing: \"Should I ask the client?\" Draft, send on yes, resolved on both laptops.",
]
text(s, 0.6, 1.85, 5.6, 5.0, demo, 14, WHITE, spacing=9)
real = [
    "Qdrant Edge memory + hybrid search",
    "Conflict detection + tier filters",
    "Offline on two physical laptops",
    "Hub sync, outbox, partial snapshots",
    "Conflict ownership + resolution",
    "Gmail, WhatsApp export, typed notes",
]
mock = [
    "Phone app",
    "Live WhatsApp stream (replayed)",
    "Screen understanding",
    "Task execution beyond reminders",
    "Long-term pattern learning",
    "Voice (stretch goal)",
]
box(s, 6.5, 1.85, 3.05, 0.5, fill=GREEN, content=[("REAL", {"size": 14, "bold": True, "color": BG})])
box(s, 9.65, 1.85, 3.05, 0.5, fill=MUTED, content=[("MOCKED", {"size": 14, "bold": True, "color": BG})])
for i, (r, m) in enumerate(zip(real, mock)):
    y = 2.45 + i * 0.66
    box(s, 6.5, y, 3.05, 0.58, fill=PANEL, content=r, size=12, color=WHITE)
    box(s, 9.65, y, 3.05, 0.58, fill=PANEL, content=m, size=12, color=MUTED)

# ---------------------------------------------------------------- 10. Business + close
s = new_slide(10)
header(s, "Business", "Small teams, priced per team per month", GREEN)
biz = [
    (AMBER, "Who", "Small field teams: film crews, event teams, agencies, site crews"),
    (GREEN, "Model", "Per team, per month. One small home hub per team."),
    (INDIGO, "Customer zero", "Us: a 4-person media team with frequent no-signal shoots"),
]
for i, (c, t, d) in enumerate(biz):
    x = 0.6 + i * 4.1
    box(s, x, 1.9, 3.85, 1.5, fill=PANEL, line=c, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, content=[
        (t, {"size": 18, "bold": True, "color": c}),
        (d, {"size": 15, "color": MUTED}),
    ])
text(s, 0.6, 4.0, 12.1, 1.3,
     "Shared memory for small teams that work where the signal doesn't.",
     26, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 5.2, 12.1, 0.5,
     "The same layer works for any device that must remember offline.",
     18, MUTED, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.9, 12.1, 0.5, "github.com/ImTani/quorum-edge-memory", 15, INDIGO, font=MONO,
     align=PP_ALIGN.CENTER)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Quorum-Pitch.pptx")
prs.save(out)
print("saved", out)
